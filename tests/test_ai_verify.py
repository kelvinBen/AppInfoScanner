#! /usr/bin/python3
# -*- coding: utf-8 -*-
# Author: kelvinBen (微信/WeChat: bromomo )
# Github: https://github.com/kelvinBen/AppInfoScanner
# Gitee: https://gitee.com/kelvin_ben/AppInfoScanner
import os
import unittest
from unittest import mock

from tests import load_runtime_config
load_runtime_config()
import libs.core as cores
from libs.core import ai_verify
from libs.core import default_config as dc


def fake_resp(payload, status=200):
    resp = mock.MagicMock()
    resp.status_code = status
    resp.raise_for_status.return_value = None
    resp.json.return_value = payload
    return resp


PROVIDERS = {
    "cloud-deepseek": {
        "kind": "openai_compatible",
        "base_url": "https://api.deepseek.com/v1",
        "api_key": "${AIS_TEST_KEY}",
        "model_text": "deepseek-chat",
        "model_vision": "",
    },
    "local-ollama": {
        "kind": "openai_compatible",
        "base_url": "http://127.0.0.1:11434/v1",
        "api_key": "",
        "model_text": "qwen2.5-coder:7b",
        "model_vision": "qwen2.5-vl:7b",
    },
    "cloud-claude": {
        "kind": "anthropic",
        "base_url": "https://api.anthropic.com",
        "api_key": "sk-test",
        "model_text": "claude-sonnet-4",
    },
}


class TestExpandEnv(unittest.TestCase):
    def test_expand_defined(self):
        os.environ["AIS_TMP_X"] = "v1"
        self.assertEqual(ai_verify.expand_env("k=${AIS_TMP_X}"), "k=v1")

    def test_undefined_kept(self):
        self.assertEqual(ai_verify.expand_env("${AIS_NO_SUCH_VAR_42}"), "${AIS_NO_SUCH_VAR_42}")

    def test_none(self):
        self.assertEqual(ai_verify.expand_env(None), "")


class TestProviderOrder(unittest.TestCase):
    # 13-1: 本地 base_url 默认最高优先(数据不出内网), 显式 priority 尊重
    def with_providers(self, providers):
        cfg = {"provider": providers}
        with mock.patch.object(cores.config, "ai_verify", cfg):
            return ai_verify.providers()

    def test_local_first(self):
        ordered = self.with_providers(PROVIDERS)
        self.assertEqual(ordered[0][0], "local-ollama")   # 本地默认 prio 0

    def test_explicit_priority(self):
        provs = dict(PROVIDERS)
        provs["cloud-claude"] = dict(provs["cloud-claude"], priority=-1)
        ordered = self.with_providers(provs)
        self.assertEqual(ordered[0][0], "cloud-claude")

    def test_missing_base_url_skipped(self):
        provs = dict(PROVIDERS)
        provs["broken"] = {"kind": "openai_compatible"}
        names = [n for n, _ in self.with_providers(provs)]
        self.assertNotIn("broken", names)

    def test_capability_slot(self):
        with mock.patch.object(cores.config, "ai_verify", {"provider": PROVIDERS}):
            name, p = ai_verify.select("vision")
            self.assertEqual(name, "local-ollama")  # cloud 无 vision 槽被跳过


class TestChatProtocols(unittest.TestCase):
    def test_openai_compatible_shape(self):
        p = PROVIDERS["cloud-deepseek"]
        os.environ["AIS_TEST_KEY"] = "sk-real"
        with mock.patch.object(ai_verify.requests, "post",
                               return_value=fake_resp({"choices": [{"message": {"content": "ok"}}]})) as post:
            out = ai_verify.chat(p, "hi")
        self.assertEqual(out, "ok")
        url, kwargs = post.call_args[0][0], post.call_args[1]
        self.assertEqual(url, "https://api.deepseek.com/v1/chat/completions")
        self.assertEqual(kwargs["headers"]["Authorization"], "Bearer sk-real")
        self.assertEqual(kwargs["json"]["model"], "deepseek-chat")

    def test_anthropic_shape(self):
        p = PROVIDERS["cloud-claude"]
        with mock.patch.object(ai_verify.requests, "post",
                               return_value=fake_resp({"content": [{"type": "text", "text": "claude-ok"}]})) as post:
            out = ai_verify.chat(p, "hi")
        self.assertEqual(out, "claude-ok")
        url, kwargs = post.call_args[0][0], post.call_args[1]
        self.assertEqual(url, "https://api.anthropic.com/v1/messages")
        self.assertEqual(kwargs["headers"]["x-api-key"], "sk-test")
        self.assertEqual(kwargs["headers"]["anthropic-version"], "2023-06-01")

    def test_vision_slot_used(self):
        p = PROVIDERS["local-ollama"]
        with mock.patch.object(ai_verify.requests, "post",
                               return_value=fake_resp({"choices": [{"message": {"content": "v"}}]})) as post:
            ai_verify.chat(p, "img", capability="vision")
        self.assertEqual(post.call_args[1]["json"]["model"], "qwen2.5-vl:7b")


class TestFallback(unittest.TestCase):
    # 13-1: 懒降级——按优先级逐个尝试, 失败换下一个
    def test_first_fail_second_ok(self):
        # 显式优先级稳定顺序: local(0) -> deepseek(10) -> claude(20)
        provs = {k: dict(v, **{"priority": {"local-ollama": 0, "cloud-deepseek": 10,
                                            "cloud-claude": 20}[k]})
                 for k, v in PROVIDERS.items()}
        good = fake_resp({"choices": [{"message": {"content": "from-deepseek"}}]})
        with mock.patch.object(cores.config, "ai_verify", {"provider": provs}), \
                mock.patch.object(ai_verify.requests, "post",
                                  side_effect=[Exception("local down"), good]) as post:
            name, out = ai_verify.call_with_fallback("hi")
        self.assertEqual((name, out), ("cloud-deepseek", "from-deepseek"))
        self.assertEqual(post.call_count, 2)

    def test_all_fail(self):
        with mock.patch.object(cores.config, "ai_verify", {"provider": PROVIDERS}), \
                mock.patch.object(ai_verify.requests, "post", side_effect=Exception("net")):
            self.assertEqual(ai_verify.call_with_fallback("hi"), (None, None))


class TestAiVerifyConfig(unittest.TestCase):
    def test_defaults(self):
        self.assertFalse(dc.DEFAULTS["ai_verify"]["enabled"])
        self.assertEqual(dc.DEFAULTS["ai_verify"]["provider"], {})
        self.assertIn("ai_verify", dc.DEEP_MERGE_DICT_KEYS)

    def test_mixed_table_serialization_roundtrip(self):
        import tomllib
        dumped = dc.dump_config(dc.DEFAULTS)
        self.assertIn("[ai_verify]", dumped)
        self.assertIn("[ai_verify.provider]", dumped)
        self.assertIn("enabled = false", dumped)
        parsed = tomllib.loads(dumped)
        self.assertEqual(parsed["ai_verify"]["granularity"], "method")
        self.assertEqual(parsed["ai_verify"]["verify_buckets"], ["pii", "credentials"])

    def test_user_provider_deep_merge(self):
        import tomllib
        user = tomllib.loads("""
[ai_verify]
enabled = true

[ai_verify.provider.my-ollama]
kind = "openai_compatible"
base_url = "http://127.0.0.1:11434/v1"
model_text = "qwen2.5:7b"
""")
        merged = dc.merge_over_defaults(user)
        self.assertTrue(merged["ai_verify"]["enabled"])
        self.assertEqual(merged["ai_verify"]["provider"]["my-ollama"]["model_text"], "qwen2.5:7b")
        # 用户未动的默认键保留
        self.assertEqual(merged["ai_verify"]["granularity"], "method")


class TestCollectTargets(unittest.TestCase):
    # 13-1b: 只收 pii/credentials 两桶, other/url 不送
    def test_bucket_filter(self):
        from libs.core import ai_verify
        rd = {"/tmp/a.js": ["[Phone_CN]-->: 13912345678",
                            "[Authorization_Header]-->: Bearer x",
                            "[Suspected_SecretKey]-->: abcdef0123456789",
                            "https://x.cn/1"]}
        targets = ai_verify.collect_targets(rd)
        self.assertIn("/tmp/a.js", targets)
        names = [n for n, _ in targets["/tmp/a.js"]]
        self.assertIn("Phone_CN", names)               # pii
        self.assertIn("Authorization_Header", names)   # credentials
        self.assertNotIn("Suspected_SecretKey", names) # other 桶不送

    def test_disabled_returns_empty(self):
        from libs.core import ai_verify
        with mock.patch.object(cores.config, "ai_verify", {"enabled": False, "provider": {}}):
            self.assertEqual(ai_verify.run_verify({"/a": ["[Phone_CN]-->: 13912345678"]}), [])


class TestExtractRegions(unittest.TestCase):
    def test_method_granularity_smali_cut(self):
        import tempfile, os
        from libs.core import ai_verify
        tmp = tempfile.mkdtemp()
        fp = os.path.join(tmp, "A.smali")
        with open(fp, "w") as f:
            f.write("junk-before\n.method public getIdCard()Ljava/lang/String;\n"
                    "const-string v0, \"330227199303025415\"\n"
                    ".end method\njunk-after")
        region = ai_verify.extract_regions(fp, ["330227199303025415"], "method", 6000)
        self.assertIn(".method", region)
        self.assertIn("330227199303025415", region)
        self.assertNotIn("junk-before", region)

    def test_file_granularity_truncated(self):
        import tempfile, os
        from libs.core import ai_verify
        tmp = tempfile.mkdtemp()
        fp = os.path.join(tmp, "big.js")
        with open(fp, "w") as f:
            f.write("x" * 50000)
        region = ai_verify.extract_regions(fp, [], "file", 1000)
        self.assertLessEqual(len(region), 3000)   # max_chars = tokens*3


class TestParseVerdicts(unittest.TestCase):
    def test_fenced_json(self):
        from libs.core import ai_verify
        text = ('结论如下:\n```json\n[{"value": "13912345678", "verdict": "confirm", '
                '"type": "手机号", "basis": "getIdCard() 返回值", "call_chain": "", "flows": []}]\n```')
        out = ai_verify.parse_verdicts(text)
        self.assertEqual(len(out), 1)
        self.assertEqual(out[0]["verdict"], "confirm")
        self.assertEqual(out[0]["type"], "手机号")

    def test_garbage_returns_empty(self):
        from libs.core import ai_verify
        self.assertEqual(ai_verify.parse_verdicts("模型拒绝回答"), [])
        self.assertEqual(ai_verify.parse_verdicts(None), [])


class TestRunVerifyPipeline(unittest.TestCase):
    def test_enabled_pipeline(self):
        import tempfile, os
        from libs.core import ai_verify
        tmp = tempfile.mkdtemp()
        fp = os.path.join(tmp, "a.smali")
        with open(fp, "w") as f:
            f.write('.method public getIdCard()\nconst-string v0, "330227199303025415"\n.end method')
        rd = {fp: ["[IDCard_CN_18]-->: 330227199303025415"]}
        verdict_json = ('[{"value":"330227199303025415","verdict":"confirm","type":"身份证号",'
                        '"basis":"getIdCard() 硬编码返回","call_chain":"getIdCard()->registerAccountProvider",'
                        '"flows":["test"]}]')
        cfg = {"enabled": True, "provider": {"p1": {"kind": "openai_compatible",
               "base_url": "http://127.0.0.1:11434/v1", "model_text": "m"}},
               "verify_buckets": ["pii", "credentials"], "granularity": "method",
               "max_file_tokens": 6000, "batch_size": 5, "timeout": 30}
        with mock.patch.object(cores.config, "ai_verify", cfg), \
                mock.patch.object(ai_verify, "call_with_fallback", return_value=("p1", verdict_json)):
            out = ai_verify.run_verify(rd)
        self.assertEqual(len(out), 1)
        self.assertEqual(out[0]["type"], "身份证号")
        self.assertEqual(out[0]["flows"], ["test"])


class TestAiVerdictReport(unittest.TestCase):
    def test_txt_and_xlsx_carry_verdicts(self):
        import tempfile, os, json
        from libs.core import report
        tmp = tempfile.mkdtemp()
        data = report.classify_results({}, set(), set())
        meta = report.build_report_meta("t.apk", "Android")
        extra = {"components": [], "components_detail": [], "permissions": [], "shell": [],
                 "sniffer": [], "ai_verdict": [{"value": "3302***", "verdict": "confirm",
                                                "type": "身份证号", "basis": "getIdCard()",
                                                "call_chain": "A->B", "flows": ["test"]}],
                 "details": {}}
        xp = os.path.join(tmp, "r.xlsx")
        tp = os.path.join(tmp, "r.txt")
        report.write_xlsx_report(xp, meta, data, extra)
        report.write_txt_report(tp, meta, data, extra)
        from openpyxl import load_workbook
        ws = load_workbook(xp)["AI_Verdict"]
        self.assertIn("身份证号", [c.value for c in ws[2]])
        txt = open(tp).read()
        self.assertIn("[AI 复核与调用链] (1)", txt)
        self.assertIn("A->B", txt)


if __name__ == "__main__":
    unittest.main()
