#! /usr/bin/python3
# -*- coding: utf-8 -*-
# Author: kelvinBen (微信/WeChat: bromomo )
# Github: https://github.com/kelvinBen/AppInfoScanner
# Gitee: https://gitee.com/kelvin_ben/AppInfoScanner
# AppInfoScanner v1.0.10 输出体系：结果分类聚合与 txt/json/xlsx 三格式报告。
#
# 输入是 ParsesThreads 产出的 result_dict {文件路径: set(结果字符串)}，
# 分类规则：
#   - "[规则集]-->: 值" 形态 -> 凭据(AK/SK)或个人/企业敏感信息(PII)，按
#     filter_ak_map / filter_pii_map 的规则集名归属
#   - 带 scheme 的地址(http(s)/jdbc/redis/... 同 filter_strs 规则口径) -> URL，
#     从中提取 host 聚合为域名清单
#   - 纯 IPv4/IPv6 -> IP 清单，IPv4 再分内网(10./172.16-31./192.168.与169.254链路本地)
#     与公网；私网与链路本地在渗透场景是高价值资产
#   - 其余(格式串/裸 token 等) -> other，只进明细区
import re
import json
import datetime

SNIFFER_HEADER = ("Number", "IP/URL", "Domain", "Status", "IP",
                  "Server", "Title", "CDN", "Finger")

# 与 default_config.filter_strs 的协议规则同口径：带 scheme 的视为 URL
_URL_PATTERN = re.compile(
    r'(?i)^(?:jdbc(?::[a-z0-9]+)*:(?:@)?//|(?:https?|ftps?|sftp|wss?|ssl|tcp|udp|ssh|telnet|smtp|imap|pop3?|ldaps?|rtmps?|rtsps?|mysql|mariadb|mssql|mongodb|redis|memcached|amqp|mqtt|file|gopher)://)')
_IP4_PATTERN = re.compile(r'^(?:\d{1,3}\.){3}\d{1,3}(?::\d{1,5})?$')
_SENSITIVE_PATTERN = re.compile(r'^\[([A-Za-z0-9_]+)\]-->:?\s*(.+)$', re.S)
_IPV6_HINT = re.compile(r'^[0-9a-fA-F:]+$')


# 从URL/地址串取host
def extract_host(value):
    rest = value
    if "://" in rest:
        rest = rest.split("://", 1)[1]
    elif "@//" in rest:  # jdbc:oracle:thin:@//host:port 形态
        rest = rest.rsplit("@//", 1)[1]
    rest = rest.split("/", 1)[0]
    if "@" in rest:
        rest = rest.rsplit("@", 1)[1]
    if rest.startswith("["):
        return rest[1:rest.index("]")] if "]" in rest else rest
    return rest.rsplit(":", 1)[0]


# 判定是否私网/链路本地IPv4
def is_private_ip4(ip):
    try:
        parts = [int(x) for x in ip.split(".")]
    except ValueError:
        return False
    if parts[0] == 10 or parts[0] == 192 and parts[1] == 168:
        return True
    if parts[0] == 172 and 16 <= parts[1] <= 31:
        return True
    if parts[0] == 169 and parts[1] == 254:  # 链路本地(含云元数据端点)
        return True
    return False


# 从地址串提取端口(URL/ip:port形态; 无端口或IPv6返回None)
def extract_port(value):
    rest = value
    if "://" in rest:
        rest = rest.split("://", 1)[1]
    elif "@//" in rest:
        rest = rest.rsplit("@//", 1)[1]
    rest = rest.split("/", 1)[0]
    if "@" in rest:
        rest = rest.rsplit("@", 1)[1]
    if rest.startswith("["):
        return None
    if ":" in rest:
        port = rest.rsplit(":", 1)[-1]
        if port.isdigit() and 1 <= int(port) <= 65535:
            return port
    return None


# 内网/回环/链路本地IPv4与IPv6字面量不嗅探(仅对IP字面量生效,域名照常)
def sniff_allowed(host):
    host = (host or "").strip().strip("[]").lower()
    if not host:
        return False
    if host == "localhost":  # 回环域名，不打到分析机本机
        return False
    if ":" in host:  # IPv6 字面量(含 ::1/fe80/fc00 ULA)，一律不嗅探
        return False
    parts = host.split(".")
    if len(parts) == 4 and all(p.isdigit() and 0 <= int(p) <= 255 for p in parts):
        first, second = int(parts[0]), int(parts[1])
        if first in (10, 127, 0) or (first == 192 and second == 168):
            return False
        if first == 172 and 16 <= second <= 31:
            return False
        if first == 169 and second == 254:  # 链路本地(含云元数据端点)
            return False
    return True


# 把result_dict分类聚合为报告数据结构
def classify_results(result_dict, ak_names, pii_names):
    data = {
        "urls": {},
        "hosts": {},
        "ip_private": {},
        "ip_public": {},
        "ip_v6": {},
        "loopback": {},
        "credentials": {},
        "pii": {},
        "other": {},
    }

    def put(bucket, key, file_path):
        bucket.setdefault(key, [])
        if file_path not in bucket[key]:
            bucket[key].append(file_path)

    for file_path, results in result_dict.items():
        for value in results:
            value = value.strip()
            if not value:
                continue
            sensitive = _SENSITIVE_PATTERN.match(value)
            if sensitive:
                name, raw = sensitive.group(1), sensitive.group(2).strip()
                if name in ak_names:
                    put(data["credentials"], "%s|%s" % (name, raw), file_path)
                elif name in pii_names:
                    put(data["pii"], "%s|%s" % (name, raw), file_path)
                else:
                    put(data["other"], value, file_path)
                continue
            host = extract_host(value) if _URL_PATTERN.match(value) or _IP4_PATTERN.match(value) else ""
            # 回环(127.x / localhost)是本地服务通讯情报: 独立归类保留端口，不进常规资产桶
            if host and (host == "localhost" or host.startswith("127.")):
                put(data["loopback"], value, file_path)
                continue
            if _URL_PATTERN.match(value):
                put(data["urls"], value, file_path)
                if host:
                    put(data["hosts"], host, file_path)
                continue
            if _IP4_PATTERN.match(value):
                # 值可能带端口(独立IP规则捕获 ip:port)，按 host 部分判定内外网
                bucket = "ip_private" if is_private_ip4(host) else "ip_public"
                put(data[bucket], value, file_path)
                continue
            if ":" in value and _IPV6_HINT.match(value):
                put(data["ip_v6"], value, file_path)
                continue
            put(data["other"], value, file_path)
    for bucket in ("urls", "hosts", "ip_private", "ip_public", "ip_v6", "loopback", "other"):
        data[bucket] = dict(sorted(data[bucket].items()))
    return data


def build_report_meta(sample, task_type, tool_version="v1.0.10"):
    return {
        "sample": sample,
        "task_type": task_type,
        "tool_version": tool_version,
        "create_time": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }


def _split_sensitive(key):
    name, _, value = key.partition("|")
    return name, value


# 结构化JSON报告
def write_json_report(path, meta, data, extra):
    payload = {"meta": meta}
    for section in ("urls", "hosts", "ip_private", "ip_public", "ip_v6", "loopback"):
        payload[section] = [
            {"value": key, "files": files} for key, files in data[section].items()]
    payload["loopback"] = [
        {"value": key, "port": extract_port(key), "files": files}
        for key, files in data["loopback"].items()]
    payload["credentials"] = _group_by_rule(data["credentials"])
    payload["pii"] = _group_by_rule(data["pii"])
    payload["other"] = [
        {"value": key, "files": files} for key, files in data["other"].items()]
    payload.update(extra)  # components/permissions/shell/sniffer 等
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)


def _group_by_rule(flat):
    grouped = {}
    for key, files in flat.items():
        name, value = _split_sensitive(key)
        grouped.setdefault(name, []).append({"value": value, "files": files})
    return grouped


# TXT分区汇总报告
def write_txt_report(path, meta, data, extra):
    lines = []
    lines.append("=" * 64)
    lines.append("AppInfoScanner 扫描报告")
    lines.append("=" * 64)
    lines.append("样本: %s    类型: %s    时间: %s" % (
        meta["sample"], meta["task_type"], meta["create_time"]))
    lines.append("")

    lines.append("[概要]")
    lines.append("  URL: %d    域名: %d    内网IP: %d    公网IP: %d    IPv6: %d" % (
        len(data["urls"]), len(data["hosts"]), len(data["ip_private"]),
        len(data["ip_public"]), len(data["ip_v6"])))
    lines.append("  敏感凭据(AK/SK): %d    个人/企业敏感信息(PII): %d    其他命中: %d" % (
        len(data["credentials"]), len(data["pii"]), len(data["other"])))
    lines.append("")

    shell = extra.get("shell") or []
    lines.append("[壳检测]")
    lines.extend("  " + s for s in shell) if shell else lines.append("  未检出加固壳")
    lines.append("")

    comps = extra.get("components") or []
    lines.append("[组件]")
    lines.extend("  " + c for c in comps) if comps else lines.append("  无")
    lines.append("")

    perms = extra.get("permissions") or []
    lines.append("[敏感权限]")
    lines.extend("  " + p for p in perms) if perms else lines.append("  无")
    lines.append("")

    def dump_list(title, bucket):
        lines.append("[%s] (%d)" % (title, len(bucket)))
        for key in bucket:
            lines.append("  %s    <- %d 个文件" % (key, len(bucket[key])))
        lines.append("")

    dump_list("域名清单", data["hosts"])
    dump_list("内网IP", data["ip_private"])
    dump_list("公网IP", data["ip_public"])
    dump_list("IPv6", data["ip_v6"])
    dump_list("本地回环服务(loopback)", data["loopback"])
    loop_ports = sorted({p for p in (extract_port(k) for k in data["loopback"]) if p})
    if loop_ports:
        lines.append("  # 抓包建议(金融/股票类APP常用本地回环通讯): 宿主机起代理监听对应端口后执行:")
        for port in loop_ports:
            lines.append("  #   adb reverse tcp:{0} tcp:{0}   (设备 127.0.0.1:{0} -> 宿主机代理)".format(port))
        lines.append("")

    lines.append("[敏感凭据 AK/SK] (%d)" % len(data["credentials"]))
    for key, files in data["credentials"].items():
        name, value = _split_sensitive(key)
        lines.append("  [%s] %s" % (name, value))
    lines.append("")
    lines.append("[个人/企业敏感信息 PII] (%d)" % len(data["pii"]))
    for key, files in data["pii"].items():
        name, value = _split_sensitive(key)
        lines.append("  [%s] %s" % (name, value))
    lines.append("")

    if extra.get("sniffer"):
        lines.append("[网络嗅探] (%d)" % len(extra["sniffer"]))
        for row in extra["sniffer"][:50]:
            lines.append("  " + " | ".join(str(x) for x in row))
        lines.append("")

    lines.append("[明细] (按文件)")
    for file_path, results in sorted(extra.get("details", {}).items()):
        lines.append("  %s:" % file_path)
        for value in sorted(results):
            lines.append("      %s" % value)
    lines.append("")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))


# XLSX多sheet报告
def write_xlsx_report(path, meta, data, extra):
    import openpyxl
    workbook = openpyxl.Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary_rows = [
        ("样本", meta["sample"]), ("类型", meta["task_type"]),
        ("时间", meta["create_time"]), ("工具版本", meta["tool_version"]),
        ("URL", len(data["urls"])), ("域名", len(data["hosts"])),
        ("内网IP", len(data["ip_private"])), ("公网IP", len(data["ip_public"])),
        ("IPv6", len(data["ip_v6"])), ("敏感凭据", len(data["credentials"])),
        ("个人/企业敏感信息", len(data["pii"])),
    ]
    for row in summary_rows:
        summary.append(row)

    def sheet(name, header, rows):
        ws = workbook.create_sheet(name)
        ws.append(header)
        for row in rows:
            ws.append(row)

    sheet("Hosts", ("Host", "Files"),
          [(k, len(v)) for k, v in data["hosts"].items()])
    sheet("IPs", ("IP", "Category", "Files"),
          [(k, "private", len(v)) for k, v in data["ip_private"].items()] +
          [(k, "public", len(v)) for k, v in data["ip_public"].items()] +
          [(k, "ipv6", len(v)) for k, v in data["ip_v6"].items()])
    sheet("Loopback", ("Loopback Service", "Port", "adb reverse", "Files"),
          [(k, extract_port(k) or "-",
            "adb reverse tcp:{0} tcp:{0}".format(extract_port(k)) if extract_port(k) else "-",
            len(v)) for k, v in data["loopback"].items()])
    sheet("URLs", ("URL", "Files"),
          [(k, len(v)) for k, v in data["urls"].items()])
    _sensitive_sheet(workbook, "Credentials", data["credentials"])
    _sensitive_sheet(workbook, "PII", data["pii"])
    sheet("Components", ("Component",), [(c,) for c in extra.get("components") or []])
    sheet("Permissions", ("Permission",), [(p,) for p in extra.get("permissions") or []])
    sheet("Shell", ("Conclusion",), [(s,) for s in extra.get("shell") or []])
    if extra.get("sniffer"):
        sheet("Sniffer", SNIFFER_HEADER, [tuple(r) for r in extra["sniffer"]])
    details = extra.get("details") or {}
    sheet("Details", ("File", "Result"),
          [(fp, v) for fp, values in sorted(details.items()) for v in sorted(values)])
    workbook.save(path)


def _sensitive_sheet(workbook, name, flat):
    ws = workbook.create_sheet(name)
    ws.append(("Rule", "Value", "Files"))
    for key, files in flat.items():
        rule, value = _split_sensitive(key)
        ws.append((rule, value, len(files)))
